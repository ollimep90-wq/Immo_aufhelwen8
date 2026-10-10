"""Exportiert den Risiko-Fragebogen aus dem Excel-Prototyp nach data/fragen_voll.json.

Aufruf: python3 -I tools/export_fragen.py <prototyp.xlsx> data/fragen_voll.json
Die Excel-Datei liegt nicht im Repo (Masterarbeit des Nutzers). Das JSON ist der Stand,
mit dem die App arbeitet; Herkunft je Frage steht im Feld "quelle".
Gewichte: aus der Kopfzeile des jeweiligen Blocks bzw. Unterblocks (Absicht des Prototyps),
nicht aus den Fragezeilen (dort greift nur die erste Spalte, siehe Prototyp_Analyse.md Befund 16).
"""
import json, sys, warnings
import openpyxl
warnings.filterwarnings("ignore")

ARTEN = ["bu", "grundfaehigkeit", "unfall", "risikoleben", "pkv", "krankenzusatz", "zahnzusatz",
         "rechtsschutz", "hausrat", "wohngebaeude", "tierkranken", "tier_op"]
FILTER = {29: "standard", 30: "kinder", 31: "tier", 32: "partner", 33: "eigentum"}

src, ziel = sys.argv[1], sys.argv[2]
ws = openpyxl.load_workbook(src, data_only=True)["Insurances"]
bloecke, fragen = [], []
kopf = None  # aktuelle Gewichte aus Block-/Unterblock-Kopfzeile
for r in range(2, ws.max_row + 1):
    nr, text = ws.cell(r, 1).value, ws.cell(r, 2).value
    if not nr:
        continue
    gew = [ws.cell(r, c).value for c in range(37, 49)]
    typ = ws.cell(r, 26).value
    if not typ:  # Kopfzeile
        teile = str(nr).split(".")
        if teile[1:] == ["0", "0", "0"]:
            titel = text.split("–")[-1].split(" - ")[-1].strip()
            titel = titel.split(" (sehr wichtig")[0].replace(" (unabhängig vom „stark sein“)", "")
            bloecke.append({"nr": int(teile[0]), "titel": titel})
        if any(gew):
            kopf = {a: int(g) for a, g in zip(ARTEN, gew) if g}
        continue
    filt = [FILTER[c] for c in FILTER if ws.cell(r, c).value == 1]
    zeile = {a: int(g) for a, g in zip(ARTEN, gew) if g}
    # Unterblöcke ohne eigene Kopfgewichte (Haustiere): dann gelten die Zeilengewichte
    gewichte = zeile if any(a not in (kopf or {}) for a in zeile) else dict(kopf or {})
    fragen.append({
        "id": "p" + str(nr).replace(".", "_"), "nr": nr, "block": int(str(nr).split(".")[0]),
        "text": text.strip(), "skala": "zustimmung" if typ == "b9" else "betroffenheit",
        "filter": [f for f in filt if f != "standard"],
        "gewichte": gewichte, "quelle": f"Prototyp {nr}"
    })
# Ergänzung, nicht im Prototyp: Haftpflicht fehlt dort (Prototyp_Analyse.md Befund 12)
ARTEN.append("haftpflicht")
bloecke.append({"nr": 10, "titel": "Schäden an anderen"})
fragen.append({"id": "n10_1", "nr": "10.0.0.1", "block": 10, "skala": "betroffenheit", "filter": [],
               "text": "Stell dir vor, du verursachst aus Versehen einen Unfall, bei dem jemand schwer verletzt wird, und sollst Schadenersatz zahlen. Wie hart würde dich das treffen?",
               "gewichte": {"haftpflicht": 3}, "quelle": "neu"})
json.dump({"stand": "Excel-Prototyp Januar 2026, exportiert 10.10.2026",
           "arten": ARTEN, "bloecke": bloecke, "fragen": fragen}, open(ziel, "w"), ensure_ascii=False, indent=1)
print(len(bloecke), "Blöcke,", len(fragen), "Fragen")
